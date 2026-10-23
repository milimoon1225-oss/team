from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Proyecto-ecommerce-Creativa-APA7.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = Inches(.5)
section.footer_distance = Inches(.5)

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    pf = style.paragraph_format
    pf.line_spacing = 2
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = Inches(.5) if name == 'Normal' else Inches(0)
    pf.widow_control = True
for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 2'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Heading 3'].font.italic = True

normal = doc.styles['Normal']
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
settings = doc.settings.element
lang = OxmlElement('w:themeFontLang'); lang.set(qn('w:val'), 'es-PE'); settings.append(lang)
doc.core_properties.title = 'Proyecto de comercio electrónico para Creativa Impresos'
doc.core_properties.subject = 'Proyecto académico y propuesta de implementación de una tienda de impresión personalizada'
doc.core_properties.author = ''
doc.core_properties.keywords = 'comercio electrónico, imprenta, Creativa Impresos, APA 7'

hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.first_line_indent = Inches(0)
hp.paragraph_format.line_spacing = 1
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE')
r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '1'; r.append(t); f.append(r); hp._p.append(f)

def p(text='', indent=True, bold=False, italic=False, center=False, keep=False):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Inches(.5) if indent else Inches(0)
    if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = keep
    run = para.add_run(text); run.bold = bold; run.italic = italic
    return para

def h(text, level=2):
    return doc.add_paragraph(text, style=f'Heading {level}')

def table(number, title, heads, rows, widths, note):
    p(f'Tabla {number}', indent=False, bold=True, keep=True)
    p(title, indent=False, italic=True, keep=True)
    tb = doc.add_table(rows=1, cols=len(heads)); tb.alignment=WD_TABLE_ALIGNMENT.CENTER; tb.autofit=False
    for c,w in zip(tb.columns,widths): c.width=Inches(w)
    for c,txt in zip(tb.rows[0].cells,heads): c.text=txt
    for row in rows:
        for c,txt in zip(tb.add_row().cells,row): c.text=txt
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr()
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        avoid=OxmlElement('w:cantSplit'); trpr.append(avoid)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right','insideH','insideV']:
                bd=OxmlElement('w:'+edge)
                visible=(edge=='top' and ri==0) or (edge=='bottom' and (ri==0 or ri==len(tb.rows)-1))
                bd.set(qn('w:val'),'single' if visible else 'nil')
                if visible: bd.set(qn('w:sz'),'6'); bd.set(qn('w:color'),'000000')
                borders.append(bd)
            tcpr.append(borders)
            mar=OxmlElement('w:tcMar')
            for side,val in [('top','90'),('bottom','90'),('left','85'),('right','85')]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');mar.append(m)
            tcpr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.first_line_indent=Inches(0)
                para.paragraph_format.line_spacing=1.15
                para.paragraph_format.space_before=para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.keep_with_next=(ri < len(tb.rows)-1)
                for run in para.runs:
                    run.font.size=Pt(11); run.bold=(ri==0)
    np=p(indent=False)
    np.add_run('Nota. ').italic=True
    np.add_run(note)

title='Proyecto de comercio electrónico para Creativa Impresos'
for _ in range(3): p('',indent=False)
doc.add_paragraph(title,style='Title')
p('',indent=False)
for txt in ['Leydy Milagros Quispe Cornejo, Johan Israel Apaza Flores y Rainer Alex Cuno Enriquez',
            'Universidad Nacional de Juliaca',
            'Negocios Electrónicos', 'Docente: Roel Dante Gomez Apaza',
            'Fecha de entrega: ____________________']:
    p(txt,indent=False,center=True)
doc.add_page_break()
doc.add_paragraph(title,style='Title')
p('El proyecto plantea una tienda virtual de productos gráficos personalizados conectada con el sistema web de gestión de Creativa Impresos. Su propósito es facilitar la compra y coordinar de forma ordenada la revisión del diseño, el cobro, la producción y la entrega. La propuesta combina una estructura académica con decisiones que deberán validarse antes de operar el canal comercial.')
p('El alcance inicial comprende un catálogo acotado, compra directa para configuraciones con precio definido y cotización para trabajos especiales. La selección de productos, la cobertura, los costos y los plazos se establecerán con información comercial real. El estudio de mercado se presenta como una investigación por ejecutar, sin atribuir resultados a encuestas que todavía no se han realizado.')

h('I Presentación y análisis del negocio',1)
h('1 Resumen ejecutivo')
p('El proyecto propone desarrollar una plataforma de comercio electrónico para Creativa Impresos, orientada a la comercialización de productos gráficos personalizados y servicios relacionados con la impresión. La tienda permitirá consultar productos, seleccionar características, cargar archivos, realizar pagos y conocer el avance de los pedidos.')
p('El canal comercial se conectará con el sistema de gestión desarrollado previamente para facilitar la administración de clientes, trabajos y cobros. Se plantean dos modalidades de atención: compra directa para productos con medidas, materiales, cantidades y precios previamente definidos, y solicitud de cotización para trabajos que requieran evaluación técnica y comercial.')
p('El objetivo general es implementar un canal digital que facilite la compra y permita gestionar los pedidos desde la solicitud hasta la entrega. Los objetivos específicos son estructurar un catálogo verificable, reducir la duplicación de registros, establecer controles de aprobación y pago, y evaluar el desempeño comercial mediante indicadores de conversión, margen y recompra.')
h('2 Descripción de la empresa o proyecto')
p('Creativa Impresos es el negocio sobre el cual se desarrolla la propuesta. El proyecto parte de un sistema web con funciones de gestión de trabajos, clientes, caja, reportes y administración. El e-commerce incorporará la atención comercial dirigida a compradores externos y su vinculación con los procesos internos.')
p('Como catálogo inicial se evaluarán tarjetas de presentación, volantes, etiquetas, banners y artículos personalizados. Su incorporación dependerá de que la imprenta confirme materiales, acabados, cantidades mínimas, costos y disponibilidad. También se identificará qué trabajos se realizan internamente y cuáles requieren servicios de terceros.')
p('Cada producto contará con una ficha que explique qué incluye el precio, qué puede personalizarse, qué archivos necesita y cuánto tiempo requiere su elaboración. El catálogo publicado deberá corresponder con opciones que la imprenta pueda producir y entregar.')
h('3 Análisis del entorno y la competencia')
h('Entorno y criterios de comparación',3)
p('La impresión personalizada exige coordinar decisiones comerciales y técnicas: material, formato, cantidad, acabado, archivo y fecha de entrega. La oportunidad que se propone validar consiste en facilitar esas decisiones mediante una experiencia digital clara, acompañada por asesoría y seguimiento. La existencia de tiendas de impresión en línea constituye evidencia de oferta digital, pero no demuestra por sí sola la demanda de Creativa ni su participación potencial en el mercado.')
p('El análisis compara funciones declaradas en páginas oficiales. Se consideran claridad del catálogo, configuración, carga de archivos, revisión del diseño, compra, entrega e integración operativa. Esta revisión no constituye una prueba de calidad, satisfacción ni liderazgo de mercado. La identificación de competidores de la ciudad de atención se incorporará al trabajo de campo.')
h('Referentes nacionales',3)
p('Línea2 presenta tarjetas, volantes y productos de merchandising, con opciones y precios visibles. Su tienda indica que el importe final y la entrega se confirman por WhatsApp; por ello, muestra una combinación de catálogo digital y atención asistida (Línea2, s. f.). Para Creativa, este caso permite examinar cómo explicar las opciones y cómo comunicar cuándo el cliente tiene un precio referencial o un total confirmado.')
p('Imprime Digital describe un recorrido que incluye cargar o preparar el diseño, elegir características, pagar y recibir el pedido (Imprime Digital, s. f.). Como aprendizaje para Creativa, se propone organizar las tareas en pasos reconocibles y permitir que el cliente conozca qué información necesita antes de continuar. Ambos casos se toman como referentes nacionales; su competencia efectiva dependerá de la cobertura y del segmento atendido.')
h('Referentes competitivos internacionales',3)
p('VistaPrint es un referente comparable por la venta directa de impresos personalizados. En su oferta de volantes permite elegir características como formato, papel y cantidad, y utilizar plantillas o cargar un diseño propio (VistaPrint, s. f.). Su ayuda también explica la vista previa y la descarga de una prueba PDF del diseño (VistaPrint, 2026). La práctica transferible consiste en vincular cada opción comercial con una configuración producible y ofrecer una revisión comprensible antes de imprimir.')
p('MOO se considera un referente en tarjetas y papelería personalizada. Su catálogo presenta distintos formatos, papeles y acabados, junto con alternativas para preparar el diseño (MOO, s. f.-a). Su documentación explica la prueba PDF y advierte que no se encuentra disponible para todos los productos (MOO, s. f.-b). La aplicación propuesta para Creativa es explicar visualmente las diferencias entre materiales y registrar con claridad qué archivo y qué especificaciones revisa el comprador.')
p('Printful aporta un modelo diferente: conecta tiendas con servicios de impresión bajo demanda, preparación y envío de productos (Printful, s. f.). Se incorpora como referente operativo para analizar la conexión entre el pedido comercial y su ejecución. Esta comparación no supone que Creativa deba adoptar el mismo modelo ni que una compra deba liberarse automáticamente a producción.')
p('VistaPrint y MOO funcionan como referentes internacionales de venta de impresos personalizados; Printful permite estudiar un modelo de producción y cumplimiento de pedidos. No se atribuye a estas empresas cobertura comercial en Perú a partir de las páginas examinadas. Su condición de competidores directos en el mercado elegido requerirá verificar destinos, costos y condiciones de atención.')
table(1,'Comparación de referentes y prácticas aplicables a Creativa Impresos',
      ['Referente','Ámbito y modelo','Práctica observada','Aplicación propuesta'],[
      ['Línea2','Nacional\nVenta asistida','Catálogo y confirmación por WhatsApp','Distinguir precio referencial y confirmado'],
      ['Imprime Digital','Nacional\nImpresión en línea','Archivo, selección, pago y entrega','Ordenar el recorrido de compra'],
      ['VistaPrint','Internacional\nVenta directa','Configuración y prueba del diseño','Opciones producibles y revisión'],
      ['MOO','Internacional\nVenta directa','Materiales, acabados y prueba PDF según producto','Explicar diferencias y requisitos'],
      ['Printful','Internacional\nImpresión bajo demanda','Conexión de tienda con producción y envío','Integrar pedidos con la gestión interna']],
      [1.05,1.25,2.0,2.2],
      'Elaboración propia a partir de Línea2 (s. f.), Imprime Digital (s. f.), VistaPrint (s. f., 2026), MOO (s. f.-a, s. f.-b) y Printful (s. f.). Las aplicaciones son propuestas para Creativa, no funciones ya implementadas.')
h('Diferenciación y posibilidades de crecimiento',3)
p('La comparación sugiere que el nivel de servicio depende de la claridad del catálogo y de la coordinación entre compra, diseño y producción. Para Creativa se propone una diferenciación basada en asesoría sobre materiales, revisión técnica de archivos, aprobación identificable del arte y seguimiento del trabajo. Estas ventajas deberán contrastarse con las necesidades del cliente y con la capacidad real de la imprenta.')
p('El crecimiento se plantea en etapas. Primero se validará una familia pequeña de productos y su recorrido completo de compra; después se ampliarán las configuraciones, la recompra y las integraciones; finalmente se evaluará extender la cobertura. La adopción de prácticas internacionales eleva el estándar de la experiencia, pero la expansión geográfica requerirá comprobar márgenes, embalaje, transporte y cumplimiento de plazos.')
p('Como evaluación preliminar, el sistema interno constituye una fortaleza; la estructuración del catálogo es un aspecto por resolver; la compra y el seguimiento digital representan una oportunidad por comprobar; y los principales riesgos son competir solo por precio, aceptar una carga de trabajo excesiva o producir un archivo incorrecto.')
h('4 Estudio de mercado y segmentación')
p('El estudio buscará conocer qué productos necesitan los clientes, cómo compran actualmente y qué dificultades encuentran al contratar una imprenta. Se evaluarán tres segmentos: emprendedores y pequeños negocios que requieren material para presentar su marca; empresas e instituciones con compras recurrentes; y personas u organizadores de eventos que solicitan productos para una ocasión.')
p('La validación incluirá entrevistas exploratorias, revisión de pedidos históricos disponibles y pruebas de compra. Se propone iniciar con 15 entrevistas seleccionadas intencionalmente entre compradores potenciales y actuales, con representación de los segmentos. Este ejercicio permitirá identificar patrones y ajustar hipótesis; no se tratará como una muestra estadísticamente representativa.')
p('Se consultará qué compran, con qué frecuencia, qué presupuesto manejan, si cuentan con diseño y qué importancia asignan al precio, la asesoría y la fecha de entrega. También se observará si pueden configurar un producto, cargar un archivo y distinguir una compra de una cotización. El segmento prioritario se elegirá considerando recurrencia, margen y capacidad de atención.')
h('5 Buyer persona o perfil del cliente ideal')
p('Se plantea como hipótesis a Carla, propietaria de un pequeño negocio que necesita etiquetas y material publicitario. Realiza consultas desde su celular y dispone de poco tiempo para coordinar presencialmente. Busca conocer el precio total, elegir materiales sin dificultad y recibir su pedido en una fecha acordada.')
p('Puede contar con un logotipo, pero necesitar ayuda para preparar el archivo de impresión. Sus preocupaciones son recibir un producto distinto al solicitado, pagar cargos no informados o descubrir un error después de imprimir. La tienda deberá ofrecerle ejemplos claros, asesoría, aprobación del diseño y seguimiento. Este perfil se ajustará con los resultados de la investigación.')

h('II Estrategia y propuesta de valor',1)
h('6 Propuesta de valor')
p('La propuesta de valor es ofrecer productos impresos personalizados con opciones claras, asesoría y seguimiento desde el pedido hasta la entrega. Se apoyará en información comprensible sobre materiales, medidas y acabados; precio y condiciones confirmados antes de aceptar el trabajo; revisión técnica y aprobación del diseño; y facilidad para consultar el avance o repetir una compra.')
p('Cada promesa comercial deberá corresponder con un procedimiento que la imprenta pueda cumplir. Los plazos se comunicarán según el producto y su capacidad disponible; los servicios de diseño, acabados y entrega se identificarán expresamente dentro del alcance contratado.')
h('7 Modelo de negocio e ingresos')
p('Se propone vender a empresas y consumidores finales, con producción principalmente bajo pedido. Las fuentes de ingresos serán la venta de impresos, el diseño o adaptación de archivos, los acabados adicionales, los pedidos recurrentes y el servicio de entrega cuando corresponda. Los paquetes comerciales se definirán después de conocer las combinaciones que compran los clientes.')
p('Para fijar precios se considerarán materiales, preparación, mano de obra, desperdicio, servicios de terceros, comisiones de cobro y embalaje. La contribución de un pedido se calculará restando sus costos variables del ingreso neto correspondiente; esta contribución ayudará a cubrir gastos fijos. Los descuentos por volumen se aprobarán solo cuando preserven la rentabilidad prevista.')
p('Antes del lanzamiento se elaborará un presupuesto que contemple plataforma, mantenimiento, publicidad y atención comercial. El presupuesto distinguirá inversión inicial y gastos periódicos. La viabilidad se evaluará con costos y ventas observables, sin presentar proyecciones como resultados garantizados.')
h('8 Estrategia de marketing digital')
p('La estrategia buscará atraer a personas con una necesidad concreta de impresión. Se publicarán fotografías reales, ejemplos de uso y explicaciones sobre materiales, medidas y preparación de archivos. Las páginas de producto deberán resolver las dudas que impiden comprar y ofrecer un acceso visible a la atención comercial.')
p('Se propone iniciar campañas pequeñas en la zona atendida, dirigidas a productos cuya producción y margen ya estén comprobados. Las acciones se medirán mediante visitas, cotizaciones, compras completadas, conversión de cotizaciones, costo de adquisición, margen y recompra. Las campañas se ampliarán cuando generen pedidos que la imprenta pueda atender de forma rentable.')
p('La fidelización incluirá historial de pedidos, recuperación de especificaciones y facilidad para volver a comprar. Las comunicaciones publicitarias se enviarán únicamente con la autorización correspondiente y se diferenciarán de los avisos necesarios para gestionar un pedido.')

h('III Diseño e implementación digital',1)
h('9 Plataforma digital y arquitectura tecnológica')
p('La plataforma tendrá una tienda dirigida al cliente y un sistema de gestión utilizado por el equipo. La tienda mostrará productos, recibirá compras o cotizaciones, archivos y datos de entrega. La gestión interna administrará los trabajos, los cobros, la producción y la entrega. La integración deberá compartir información mediante operaciones autorizadas y evitar duplicaciones.')
p('La primera versión incluirá catálogo, fichas, selección de opciones, carrito, proceso de compra, formulario de cotización, carga de archivos, revisión y aprobación del diseño, seguimiento e historial. Un editor gráfico completo se evaluará en una etapa posterior, de acuerdo con su utilidad y costo. Para iniciar, se ofrecerán las alternativas de subir un diseño o solicitar el servicio de diseño.')
p('Los precios se comprobarán en el servidor y los clientes accederán únicamente a sus propios pedidos y archivos. La integración comprobará que cada solicitud de pago y cada pedido se registren una sola vez. Se mantendrán separados el estado del pago, la aprobación del arte y el avance de producción.')
p('El desarrollo se organizará en tres etapas: validar catálogo y reglas comerciales, implementar un recorrido completo y ejecutar un piloto antes de ampliar productos o automatizaciones. Los criterios de aceptación incluirán completar una compra y una cotización, corregir un archivo, registrar una aprobación y consultar el estado sin exponer información de otros clientes.')
p('El wireframe se presenta en un archivo complementario independiente, titulado Wireframe del e-commerce de Creativa Impresos. Sus pantallas representan la propuesta funcional de inicio, catálogo, ficha, cotización, carrito, compra, aprobación, seguimiento y cuenta, junto con vistas móviles. Las anotaciones explican la función de cada bloque y su relación con los procesos del proyecto.')
h('10 Procesos operativos y logística')
p('El proceso comenzará con la selección de un producto o la solicitud de una cotización. La imprenta revisará especificaciones, archivo y capacidad; confirmará importe, alcance, plazo y forma de entrega; y preparará o corregirá el diseño cuando corresponda. Después, el cliente aprobará una versión identificada del arte. La producción se liberará con esa aprobación y con el pago requerido confirmado.')
p('Antes de entregar, se verificará cantidad, medidas, impresión y acabados. El retiro o despacho quedará registrado y se ofrecerá un canal para atender incidencias. Un archivo recibido no se considerará automáticamente apto para imprimir; deberá superar la revisión técnica correspondiente.')
p('La aprobación identificará el archivo, su versión y las especificaciones. Si el cliente pide cambios después de aprobar, se revisarán nuevamente precio, archivo y plazo antes de continuar. La confirmación de compra informará las condiciones aplicables cuando el archivo deba corregirse o no pueda utilizarse.')
p('Se distinguirán los tiempos de diseño, producción y transporte. El plazo de producción comenzará cuando se cumplan las condiciones comunicadas de aprobación y pago. La fecha de entrega considerará además la capacidad disponible y, cuando existan, los servicios de terceros.')
p('La logística evaluará retiro en el local y entrega dentro de una cobertura inicial. El costo y las condiciones se mostrarán antes de confirmar la compra. La expansión dependerá de pruebas de embalaje, transporte y cumplimiento; no se ofrecerá cobertura sin comprobar que el pedido puede llegar en condiciones adecuadas.')
h('11 Pagos facturación y atención al cliente')
p('Se establecerá pago completo para los productos estandarizados cuando resulte adecuado, y adelanto con saldo para los trabajos cotizados bajo condiciones aceptadas previamente. La propuesta comercial explicará importe, vigencia, alcance, hitos y condiciones antes de que el comprador acepte.')
p('Los pagos electrónicos se confirmarán mediante el proveedor de cobros. Cuando exista verificación manual, el pedido permanecerá pendiente hasta comprobar el abono. Una captura enviada por el cliente no confirmará automáticamente un pago. El sistema registrará por separado el estado financiero y el estado del trabajo, incluyendo pagos pendientes, rechazados y devoluciones cuando correspondan.')
p('La facturación contemplará el comprobante aplicable y su relación con el pedido. La Superintendencia Nacional de Aduanas y de Administración Tributaria (SUNAT, 2025) identifica la boleta electrónica como comprobante para consumidores finales y describe las características de la factura electrónica para las operaciones correspondientes. La emisión tributaria se incorporará mediante un procedimiento o integración específico; el recibo interno de gestión no sustituirá el comprobante aplicable.')
p('La atención al cliente definirá horarios, responsables y procedimientos para consultas, correcciones, retrasos, cancelaciones y problemas de calidad. Las condiciones se mostrarán antes de la compra y contemplarán cómo resolver diferencias entre lo contratado y lo entregado.')
p('Se incorporará un Libro de Reclamaciones virtual accesible. El Instituto Nacional de Defensa de la Competencia y de la Protección de la Propiedad Intelectual (Indecopi, s. f.) establece que el proveedor debe responder reclamos y quejas en un máximo de 15 días hábiles no prorrogables. Se conservará evidencia de su registro y atención.')
h('12 Seguridad de la información y privacidad')
p('La plataforma protegerá datos de contacto, direcciones, archivos de diseño e información de los pedidos. Se contemplarán conexiones seguras, autenticación, permisos diferenciados, archivos privados y validación de formatos y tamaños. El registro de cambios y aprobaciones permitirá conocer qué versión fue autorizada y quién realizó cada acción relevante.')
p('Se realizarán copias de seguridad y pruebas de recuperación. Los pagos se gestionarán con un proveedor especializado y se evitará almacenar datos completos de tarjetas. También se definirán responsables y procedimientos para identificar y atender incidentes de seguridad.')
p('La política de privacidad explicará los datos recogidos, sus finalidades, los destinatarios y los medios para ejercer derechos. Se definirán plazos de conservación y una autorización diferenciada para comunicaciones comerciales.')
p('La implementación contemplará la Ley N.º 29733 y su reglamento, aprobado mediante el Decreto Supremo N.º 016-2024-JUS y vigente desde el 31 de marzo de 2025 (Autoridad Nacional de Protección de Datos Personales [ANPD], 2025).')

h('Referencias',1).paragraph_format.page_break_before=True
refs=[
('Autoridad Nacional de Protección de Datos Personales. (2025, 17 de noviembre). ', 'Nuevo reglamento de protección de datos personales.', ' https://www.gob.pe/institucion/anpd/campa%C3%B1as/128319-nuevo-reglamento-de-proteccion-de-datos-personales'),
('Imprime Digital. (s. f.). ', 'Imprime tus documentos en minutos.', ' Recuperado el 15 de octubre de 2026, de https://www.imprime.digital/'),
('Instituto Nacional de Defensa de la Competencia y de la Protección de la Propiedad Intelectual. (s. f.). ', 'Libro de reclamaciones.', ' https://consumidor.gob.pe/libro-de-reclamaciones/'),
('Línea2. (s. f.). ', 'Tienda Línea2.', ' Recuperado el 15 de octubre de 2026, de https://www.lineadosperu.com/tienda/'),
('MOO. (s. f.-a). ', 'Custom business cards—Design & print online.', ' Recuperado el 15 de octubre de 2026, de https://www.moo.com/us/business-cards'),
('MOO. (s. f.-b). ', 'PDF proof.', ' https://support.moo.com/hc/en-us/articles/360035605511-PDF-proof'),
('Printful. (s. f.). ', 'How does Printful work?', ' https://www.printful.com/how-printful-works'),
('Superintendencia Nacional de Aduanas y de Administración Tributaria. (2025, 25 de septiembre). ', 'Tipos de comprobante de pago.', ' https://cpe.sunat.gob.pe/informacion_general/tipos_comprobantes_pago'),
('VistaPrint. (s. f.). ', 'Flyers.', ' Recuperado el 15 de octubre de 2026, de https://www.vistaprint.com/marketing-materials/flyers'),
('VistaPrint. (2026, 16 de junio). ', '¿Cómo obtengo una vista previa de mi diseño?', ' https://www.vistaprint.com/es/atencion-al-cliente/centro-de-ayuda/000002551'),
]
for pre,work,post in refs:
    para=p(indent=False)
    para.paragraph_format.left_indent=Inches(.5)
    para.paragraph_format.first_line_indent=Inches(-.5)
    para.add_run(pre)
    para.add_run(work).italic=True
    para.add_run(post)

# Remove inherited theme fonts and title rules from the bundled Word template.
for style in doc.styles:
    for bd in list(style.element.iter(qn('w:pBdr'))):
        bd.getparent().remove(bd)
    for fonts in style.element.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if key.endswith('Theme'): del fonts.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:
            fonts.set(qn('w:'+key),'Times New Roman')
for para in doc.paragraphs:
    para.paragraph_format.line_spacing=2
    para.paragraph_format.space_before=para.paragraph_format.space_after=Pt(0)
    for bd in list(para._p.iter(qn('w:pBdr'))):
        bd.getparent().remove(bd)
    for run in para.runs:
        run.font.name='Times New Roman'
        run.font.size=Pt(12)
        run.font.color.rgb=RGBColor(0,0,0)
doc.save(OUT)
info={'file':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'references':len(refs),'words':sum(len(x.text.split()) for x in doc.paragraphs)}
(ROOT/'qa'/'documento-control.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(info,ensure_ascii=False))
