from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Proyecto-ecommerce-Creativa-APA7.docx'
OUT = ROOT / 'Proyecto-ebusiness-Creativa-APA7.docx'
doc = Document(SOURCE)

changes = [
('Proyecto de comercio electrónico para Creativa Impresos', 'Proyecto de negocios electrónicos para Creativa Impresos'),
('El proyecto plantea una tienda virtual', 'El proyecto plantea un modelo de e-business para Creativa Impresos que integra la relación con clientes y proveedores con el diseño, la producción, el abastecimiento y la administración. El sistema web de gestión ya desarrollado constituye su base. La tienda virtual será el canal de e-commerce dentro de este conjunto, conectada con los procesos que permiten cumplir cada pedido.'),
('El alcance inicial comprende un catálogo', 'El alcance comprende ventas presenciales y digitales, seguimiento comercial, órdenes de trabajo, archivos de diseño, programación de responsables, compras y materiales, cobros, entregas y reportes. Se distinguirán las funciones existentes de las integraciones y ampliaciones propuestas. El catálogo, los costos, la cobertura y los plazos se validarán con información real; el estudio de mercado se mantiene como una investigación por ejecutar.'),
('El proyecto propone desarrollar una plataforma de comercio electrónico', 'El proyecto propone integrar digitalmente las actividades de Creativa Impresos mediante un modelo de negocios electrónicos. El propósito es conectar la captación de clientes y la venta con el diseño, la producción, los materiales, los proveedores, los cobros y la entrega, conservando una trazabilidad común por trabajo.'),
('El canal comercial se conectará', 'El sistema de gestión existente se ampliará e integrará con el canal comercial. La tienda permitirá compra directa para configuraciones con precio definido y cotización para trabajos especiales. Los pedidos presenciales y digitales se administrarán con reglas comunes de revisión, asignación, aprobación, producción y cobro, identificando siempre su origen.'),
('El objetivo general es implementar un canal digital', 'El objetivo general es diseñar un modelo de e-business que coordine los procesos comerciales y operativos de la imprenta. Los objetivos específicos son integrar datos de clientes y trabajos; vincular ventas con capacidad y materiales; controlar aprobaciones, compras y pagos; y evaluar el servicio y la rentabilidad. Las mejoras se medirán con datos previos y posteriores al piloto, sin asumir resultados anticipados.'),
('Creativa Impresos es el negocio', 'Creativa Impresos es el negocio sobre el cual se desarrolla la propuesta. Cuenta con un sistema web de gestión de trabajos, clientes, caja, reportes y administración, además de funciones para diseñadores. Esta base se conectará con el e-commerce. Las compras, el control estructurado de materiales y la programación ampliada de producción se presentan como mejoras propuestas que requieren diseño y validación.'),
('Cada producto contará con una ficha', 'Cada producto contará con una ficha comercial y especificaciones operativas vinculadas: materiales, medidas, acabados, cantidades y servicios necesarios. La oferta publicada deberá corresponder con la capacidad de producir y entregar. Los responsables internos mantendrán las reglas comerciales, los datos del cliente y la información de cada orden de trabajo.'),
('La impresión personalizada exige coordinar', 'La impresión personalizada exige coordinar material, formato, cantidad, acabado, archivo y fecha de entrega. El entorno del negocio también incluye proveedores de insumos, servicios tercerizados y transportistas. La oportunidad por validar consiste en mejorar tanto la experiencia del comprador como la coordinación interna, mediante información compartida y responsabilidades claras.'),
('El análisis compara funciones declaradas', 'El análisis compara funciones declaradas en páginas oficiales: catálogo, configuración, archivos, revisión del diseño, compra, entrega e integración operativa. No permite inferir la arquitectura interna, los costos ni las prácticas de abastecimiento de los competidores. El trabajo de campo incorporará competidores locales, plazos de proveedores y alternativas de tercerización, sin atribuir resultados que no se hayan comprobado.'),
('La comparación sugiere que el nivel de servicio', 'La comparación sugiere que el nivel de servicio depende de coordinar compra, diseño y producción. Para Creativa se propone vincular la asesoría y la aprobación del arte con disponibilidad de materiales, asignación de responsables y seguimiento. Las prácticas internacionales se usarán para elevar el estándar de servicio, mientras que su aplicación se evaluará según la capacidad real de la imprenta.'),
('El crecimiento se plantea en etapas.', 'El crecimiento se plantea en etapas: ordenar datos y procesos del sistema existente, integrar el canal comercial y los controles de diseño y producción, e incorporar compras, materiales y análisis de costos. La ampliación del catálogo o la cobertura se evaluará cuando existan evidencias de capacidad, margen y cumplimiento. La integración no supone automatizar físicamente las máquinas de impresión.'),
('Como evaluación preliminar, el sistema interno', 'Como evaluación preliminar, el sistema interno constituye una fortaleza; la integración de datos comerciales y operativos es un aspecto por resolver; la trazabilidad y la recompra representan oportunidades por validar; y los riesgos incluyen datos inconsistentes, falta de materiales, resistencia al cambio y aceptar trabajos que excedan la capacidad disponible.'),
('La validación incluirá entrevistas exploratorias', 'La validación incluirá revisión de pedidos históricos disponibles, entrevistas exploratorias y pruebas de tareas. Se propone iniciar con 15 entrevistas a compradores potenciales y actuales de los segmentos elegidos. Se analizarán necesidades y hábitos, sin tratar este grupo como una muestra estadísticamente representativa. La demanda se contrastará con materiales, urgencia, tiempos y capacidad necesarios para atenderla.'),
('Se consultará qué compran, con qué frecuencia', 'Se consultará qué compran, con qué frecuencia, con qué presupuesto y qué apoyo de diseño requieren. En paralelo se observará el trabajo de recepción, diseño, producción y administración, y se revisarán tiempos de proveedores y faltantes de insumos. El segmento prioritario se elegirá considerando recurrencia, contribución económica y capacidad de atención.'),
('La propuesta de valor es ofrecer', 'La propuesta de valor es ofrecer impresos personalizados con asesoría, condiciones claras y cumplimiento trazable. La experiencia del cliente se respaldará con información compartida sobre especificaciones, arte aprobado, responsables, materiales, cobros y entrega. La integración buscará evitar recapturas y facilitar respuestas consistentes sobre cada trabajo.'),
('Cada promesa comercial deberá corresponder', 'Cada promesa comercial deberá corresponder con un procedimiento verificable. La fecha se confirmará considerando materiales, carga de trabajo y terceros. El valor para la imprenta será disponer de información para programar tareas, controlar costos, dar seguimiento a saldos y tomar decisiones; estos beneficios se comprobarán durante el piloto.'),
('Se propone vender a empresas y consumidores finales', 'El modelo combinará atención presencial y digital a empresas y consumidores finales, con producción principalmente bajo pedido. Los ingresos provendrán de impresos, diseño, acabados y servicios de entrega cuando correspondan. Las actividades clave serán vender, diseñar, abastecer, producir y atender. Los recursos comprenderán personal, equipos, materiales y sistemas; los socios incluirán proveedores, talleres externos, transportistas y proveedores tecnológicos.'),
('Antes del lanzamiento se elaborará un presupuesto', 'El presupuesto distinguirá inversión inicial, operación y capacitación. Se evaluarán plataforma, mantenimiento, publicidad, atención, integración de datos y soporte al equipo. Los posibles ahorros por reducir reprocesos o recapturas se medirán como mejoras de eficiencia. Se distinguirán valor de ventas, margen y dinero cobrado, considerando también compras y pagos pendientes.'),
('La fidelización incluirá historial de pedidos', 'La fidelización utilizará historial, especificaciones y seguimiento comercial para facilitar la recompra. Se asignará un responsable a cada oportunidad y se registrará su resultado, con comunicaciones publicitarias autorizadas. La información de campañas y consultas se contrastará con la capacidad y la disponibilidad de materiales antes de ampliar promociones.'),
('La plataforma tendrá una tienda dirigida al cliente', 'La arquitectura integrará cinco componentes: canal comercial y relación con clientes; gestión de trabajos y archivos; coordinación de producción y responsables; compras y materiales; y caja y reportes. Se aprovechará la base existente y se definirán ampliaciones por etapas. Los datos se compartirán mediante operaciones autorizadas, sin exigir aplicaciones separadas cuando el sistema actual pueda ampliarse.'),
('La primera versión incluirá catálogo, fichas', 'El núcleo integrado vinculará cliente, cotización o venta, orden de trabajo, especificaciones, archivos y versiones, responsable y fechas. Se añadirá una relación entre orden, materiales reservados o consumidos, costos y movimientos de cobro. El catálogo, el carrito y la compra formarán el componente de e-commerce. Los campos y flujos de materiales y compras deberán desarrollarse y validarse antes de su uso operativo.'),
('Los precios se comprobarán en el servidor', 'Para cada dato se establecerán origen, responsable y momento de actualización. El área comercial confirmará precio y alcance; diseño registrará versiones; producción actualizará avance y consumo; administración verificará cobros y egresos. Cada cliente accederá solo a sus pedidos. Se comprobarán precios y permisos en el servidor y se evitarán pedidos, pagos y movimientos de material duplicados.'),
('El desarrollo se organizará en tres etapas', 'La implementación tendrá tres etapas: documentar procesos, roles y datos del sistema existente; integrar ventas, diseño, aprobación, producción y entrega; y añadir compras, materiales y reportes consolidados. El piloto comprobará un pedido de principio a fin y una compra de insumos, incluyendo corrección de archivo, saldo pendiente, falta de material y recuperación de información. Se capacitará a cada rol en sus tareas.'),
('El wireframe se presenta en un archivo', 'El wireframe se presenta en un archivo complementario independiente, titulado Wireframe del e-business de Creativa Impresos. Representa las pantallas comerciales y móviles junto con la gestión de órdenes, materiales y abastecimiento, y el panel administrativo. Las anotaciones distinguen la base existente de las funciones propuestas y explican cómo se conectan los procesos.'),
('El proceso comenzará con la selección', 'El proceso comercial comenzará con una solicitud presencial o digital. Se revisarán especificaciones, archivo, capacidad y materiales antes de confirmar importe, plazo y entrega. El trabajo se asignará a un responsable, se preparará el diseño y el cliente aprobará una versión identificada. La producción se liberará con el pago requerido confirmado, el arte vigente aprobado y los recursos necesarios disponibles.'),
('Antes de entregar, se verificará cantidad', 'Producción registrará avance, materiales consumidos, incidencias y tiempos por trabajo. El control de calidad comprobará cantidad, medidas, impresión y acabados. La entrega quedará vinculada con el retiro o despacho y con la atención posterior. Un archivo recibido requerirá revisión técnica antes de considerarse apto para imprimir.'),
('Se distinguirán los tiempos de diseño', 'La programación considerará tiempos de diseño, abastecimiento, producción y transporte, además de responsables, equipos y terceros. Una falta de material o un cambio de prioridad requerirá revisar el compromiso de entrega y comunicar su impacto. El registro digital coordinará tareas humanas y recursos; cualquier automatización de equipos requerirá un proyecto específico.'),
('Los pagos electrónicos se confirmarán', 'Los cobros se confirmarán mediante el proveedor de pagos o la verificación administrativa correspondiente. Se mantendrán separados el estado financiero y el avance del trabajo. Caja relacionará adelantos, saldos, devoluciones y egresos con sus documentos de origen. Los pagos a proveedores se vincularán con la compra y recepción de insumos, evitando considerar una captura del cliente como confirmación automática.'),
('La atención al cliente definirá horarios', 'La atención al cliente definirá horarios y responsables, con acceso autorizado a arte, pagos, producción y entrega. Se establecerán procedimientos para correcciones, retrasos, cancelaciones y calidad. Administración revisará periódicamente cobros, compras, pagos y saldos para conciliar la información operativa y financiera.'),
('La plataforma protegerá datos de contacto', 'La plataforma protegerá información de clientes, personal, proveedores, diseños, costos y caja. Se establecerán accesos por función, autenticación y archivos privados. Diseño y producción recibirán los datos necesarios para su tarea; el acceso a costos y movimientos financieros se reservará a los roles autorizados. Los registros de cambios identificarán usuarios, versiones y acciones relevantes.'),
('Se realizarán copias de seguridad', 'Se realizarán copias de seguridad y pruebas de recuperación. Los pagos se gestionarán mediante proveedores especializados, evitando almacenar datos completos de tarjetas. Se designarán responsables de la información, se capacitará al personal y se definirán procedimientos para incidentes, errores de registro y continuidad del trabajo.'),
]

def rewrite(para, text):
    para.clear()
    run = para.add_run(text)
    run.font.name = 'Times New Roman'; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0,0,0)
    return para

applied = []
for prefix, text in changes:
    targets = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if not targets: raise ValueError(f'No se encontro: {prefix}')
    for para in targets: rewrite(para, text)
    applied.append({'prefix':prefix,'count':len(targets)})

def after(prefix, text):
    anchor = next(p for p in doc.paragraphs if p.text.startswith(prefix))
    node = OxmlElement('w:p'); anchor._p.addnext(node)
    para = Paragraph(node, anchor._parent); para.style = doc.styles['Normal']
    para.paragraph_format.first_line_indent = Inches(.5)
    para.paragraph_format.line_spacing = 2
    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
    rewrite(para,text)
    return para

after('El alcance comprende ventas presenciales', 'La OECD (2019) aborda el e-business mediante la integración de tecnologías en los procesos empresariales y distingue las transacciones de comercio electrónico de herramientas de gestión interna y de relación con clientes. A partir de ese marco, la propuesta de Creativa conecta su tienda con las actividades que sostienen la operación del negocio.')
after('Puede contar con un logotipo', 'Para definir los requisitos también se identificarán usuarios internos: recepción y ventas, diseño, producción y administración. Se documentará qué información utiliza cada rol, qué tareas realiza y qué decisiones autoriza. Los proveedores se analizarán como colaboradores del abastecimiento; estos perfiles complementan el buyer persona del comprador.')
after('El presupuesto distinguirá inversión inicial', 'Como ejemplo del funcionamiento propuesto, un pedido se vinculará con una orden y sus materiales; un faltante originará una solicitud de compra; la recepción actualizará existencias; y producción registrará consumo y tiempo. Al cerrar el trabajo, administración podrá comparar costo previsto y registrado, cobros y saldos. Este flujo describe integraciones por desarrollar y comprobar.')
after('La programación considerará tiempos', 'El abastecimiento comprenderá solicitud y autorización de compra, elección del proveedor, orden de compra, recepción y verificación. Se distinguirán existencias físicas, reservas y consumos para evitar descontar dos veces un material. Las mermas se registrarán por trabajo y las alertas de reposición usarán mínimos definidos por la imprenta. La coordinación con proveedores podrá iniciarse mediante registros internos, sin exigir un portal externo.')
after('La implementación tendrá tres etapas', 'Los indicadores combinarán conversión y recompra con entrega a tiempo, reprocesos, faltantes, carga por responsable, costo previsto frente al registrado y saldos pendientes. Se definirá la fórmula, el responsable y el período de cada indicador. Las metas se establecerán con una línea de base; las asignaciones y tiempos del equipo no implican incorporar una nómina completa en esta etapa.')

# Add the conceptual reference at its alphabetical position in the existing list.
reference_anchor = next(p for p in doc.paragraphs if p.text.startswith('Printful. (s. f.).'))
ref = reference_anchor.insert_paragraph_before()
ref.paragraph_format.left_indent=Inches(.5)
ref.paragraph_format.first_line_indent=Inches(-.5)
ref.paragraph_format.line_spacing=2
ref.paragraph_format.space_before=ref.paragraph_format.space_after=Pt(0)
ref.add_run('OECD. (2019). ')
ref.add_run('Measuring the digital transformation: A roadmap for the future.').italic=True
ref.add_run(' OECD Publishing. https://doi.org/10.1787/9789264311992-en')
for run in ref.runs:
    run.font.name='Times New Roman'; run.font.size=Pt(12)

doc.core_properties.title='Proyecto de negocios electrónicos para Creativa Impresos'
doc.core_properties.subject='Modelo de e-business que integra ventas gestion de trabajos produccion abastecimiento y administracion'
doc.core_properties.keywords='e-business, negocios electrónicos, imprenta, Creativa Impresos, APA 7'
doc.save(OUT)
info={'output':str(OUT),'updated_blocks':len(changes),'added_paragraphs':6,'references':11,'words':sum(len(p.text.split()) for p in doc.paragraphs)}
(ROOT/'qa'/'ebusiness-control.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(info,ensure_ascii=False))
